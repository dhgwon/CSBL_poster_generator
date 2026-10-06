# -*- coding: utf-8 -*-
"""CSBL poster layout engine (Soongsil / CSBL house format).

Everything is laid out in centimetres on a PowerPoint canvas of arbitrary size.
Each figure is placed as its own picture object, so the .pptx stays editable.

Typical use
-----------
    from poster_lib import Poster, BAR, PIC, TXT, HEAD, GRID, SPLIT

    P = Poster(w_cm=57.4, h_cm=102.0, figs="…/figs")
    P.header(title=["Line one", "Line two"], authors="A Kim¹ and B Lee¹*",
             affils=["¹Dept …", "*Corresponding author: …"], logo="…/ssu_logo.png")
    P.compose([
        ("F", [BAR("Abstract"), TXT(abstract)]),
        ("F", [BAR("Overall scheme"), PIC("fig1_scheme")]),
        ("MID", left_blocks, right_blocks),          # the one elastic region
        ("F", [BAR("Results"), GRID(["a","b","c"], [17.7]*3)]),
        ("F", [BAR("Conclusions"), TXT(*bullets, bullets=True)]),
    ])
    P.save("…/poster_v1.pptx")

`compose` gives every full-width band its natural height and hands the leftover
vertical space to the MID region, so the poster always fills the page exactly.
Watch the printed table: "*** OVERFLOW ***" means a column needs less content or
smaller figures.
"""
import os, glob
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageFont

CM = 360000                                   # EMU per centimetre
LINE = 1.21                                   # PowerPoint single-line factor


def _ttf(bold):
    pats = ["/usr/share/fonts/truetype/liberation/LiberationSans-%s.ttf",
            "/usr/share/fonts/truetype/liberation2/LiberationSans-%s.ttf",
            "/usr/share/fonts/**/DejaVuSans%s.ttf"]
    for p in pats:
        for c in glob.glob(p % ("Bold" if bold else "Regular"), recursive=True):
            return c
    raise RuntimeError("no TrueType sans font found for text measurement")


THEME = dict(accent="016694", edge="1E3664", white="FFFFFF", black="000000",
             frame_fill="F7FAFB", frame_line="B9C6CE", font="Arial")


# ----------------------------------------------------------------- block types
def BAR(text):
    """Blue section bar. Embed '\\n' for a two-line bar."""
    return {"k": "bar", "text": text}

def PIC(name, frac=1.0):
    """One figure, centred, at `frac` of the column width."""
    return {"k": "pic", "name": name, "frac": frac}

def TXT(*paras, **kw):
    """Body text. kw: size=Pt(..), bold, bullets=True, align, space_after."""
    return {"k": "txt", "paras": list(paras), **kw}

def HEAD(text, sec=True):
    """Bold sub-heading. sec=False keeps it glued to the block above."""
    return {"k": "head", "text": text, "sec": sec}

def GRID(names, widths, gap_cm=0.5, frames=None, pad_cm=0.35, lab_cm=0.85):
    """A row of figures at explicit widths (cm). `frames` = per-cell captions,
    which wraps each figure in a labelled box."""
    return {"k": "grid", "names": names, "widths": widths, "gap": int(gap_cm * CM),
            "frames": frames, "pad": int(pad_cm * CM), "lab": int(lab_cm * CM)}

def SPLIT(left, rights, wl, wr, gap_cm=0.5, vgap_cm=0.35):
    """One figure on the left, a vertical stack of figures on the right."""
    return {"k": "split", "left": left, "rights": rights, "wl": wl, "wr": wr,
            "gap": int(gap_cm * CM), "vgap": int(vgap_cm * CM)}


class Poster:
    def __init__(self, w_cm, h_cm, figs, theme=None, margin_cm=1.6, gutter_cm=1.2,
                 body_pt=None, bar_pt=None, head_pt=None, header_cm=None):
        self.W, self.H = int(w_cm * CM), int(h_cm * CM)
        self.FIGS = figs
        self.T = dict(THEME, **(theme or {}))
        self.C = {k: RGBColor.from_string(v) for k, v in self.T.items() if k != "font"}
        self.FONT = self.T["font"]

        # Type scale follows the canvas: a 90 cm print poster and a 57 cm DID
        # screen are read from the same distance relative to their own width.
        s = w_cm / 57.4
        self.BODY = body_pt or Pt(round(15 * s, 1))
        self.BAR_SZ = bar_pt or Pt(round(19 * s, 1))
        self.HEAD_SZ = head_pt or Pt(round(20 * s, 1))

        self.MARGIN, self.GUTTER = int(margin_cm * CM), int(gutter_cm * CM)
        self.COL_W = (self.W - 2 * self.MARGIN - self.GUTTER) // 2
        self.FULL_W = self.W - 2 * self.MARGIN
        self.COL_X = {"L": self.MARGIN, "R": self.MARGIN + self.COL_W + self.GUTTER,
                      "F": self.MARGIN}
        self.COL_WIDTH = {"L": self.COL_W, "R": self.COL_W, "F": self.FULL_W}

        self.BAR_H1, self.BAR_H2 = int(1.45 * CM * s), int(2.30 * CM * s)
        self.GAP_BAR, self.GAP_FIG = int(0.32 * CM), int(0.40 * CM)
        self.GAP_MIN, self.BAND_GAP = int(0.95 * CM), int(0.90 * CM)
        self.HEADER_H = int((header_cm or 9.4 * s) * CM)
        self.BOT_PAD = int(1.3 * CM)

        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = self.W, self.H
        self.slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self._fc = {}

    # ------------------------------------------------------------- text metrics
    def _font(self, pt, bold):
        k = (pt, bold)
        if k not in self._fc:
            self._fc[k] = ImageFont.truetype(_ttf(bold), int(pt * 10))
        return self._fc[k]

    def wrap_lines(self, text, pt, bold, width_emu):
        f = self._font(pt, bold)
        limit = (width_emu - 2 * 91440) / 12700.0 * 10
        lines, cur = 0, ""
        for w in text.split():
            t = w if not cur else cur + " " + w
            if f.getlength(t) <= limit:
                cur = t
            else:
                lines += 1; cur = w
        return lines + (1 if cur else 0)

    def txt_h(self, paras, w, size, bullets=False, space_after=0, bold=False):
        n = sum(self.wrap_lines(("•  " + p) if bullets else p, size.pt, bold, w)
                for p in paras)
        return int(n * size.pt * LINE * 12700) + len(paras) * space_after + 70000

    # ------------------------------------------------------------------ drawing
    def _fmt(self, run, size, bold, color):
        f = run.font
        f.name, f.size, f.bold, f.italic = self.FONT, size, bold, False
        f.color.rgb = color

    def textbox(self, x, y, w, paras, size=None, bold=False, color=None,
                align=PP_ALIGN.JUSTIFY, bullets=False, space_after=Pt(0)):
        size = size or self.BODY
        color = color if color is not None else self.C["black"]
        tb = self.slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)),
                                           Emu(300000))
        tf = tb.text_frame; tf.word_wrap = True
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = align; para.space_after = space_after
            r = para.add_run(); r.text = ("•  " + p) if bullets else p
            self._fmt(r, size, bold, color)
        return tb

    def _asp(self, name):
        iw, ih = Image.open(f"{self.FIGS}/{name}.png").size
        return iw / ih

    # ------------------------------------------------------------ measuring
    def _grid_layout(self, b, cw):
        fr = b.get("frames")
        pad = b.get("pad", 0) if fr else 0
        lab = b.get("lab", 0) if fr else 0
        widths = [int(v * CM) for v in b["widths"]]
        row_w = sum(widths) + b["gap"] * (len(widths) - 1) + (2 * pad * len(widths))
        x = (cw - row_w) // 2
        placed, rh = [], 0
        for i, (n, w) in enumerate(zip(b["names"], widths)):
            h = int(w / self._asp(n))
            placed.append((n, x + pad, lab + pad, w, h))
            if fr:
                placed[-1] = placed[-1] + (x, w + 2 * pad, h + lab + 2 * pad, fr[i])
            x += w + b["gap"] + 2 * pad
            rh = max(rh, h + lab + 2 * pad if fr else h)
        return rh, placed

    def _split_layout(self, b, cw):
        wl, wr = int(b["wl"] * CM), int(b["wr"] * CM)
        hl = int(wl / self._asp(b["left"]))
        hs = [int(wr / self._asp(n)) for n in b["rights"]]
        hr = sum(hs) + b["vgap"] * (len(hs) - 1)
        rh = max(hl, hr)
        x0 = (cw - (wl + b["gap"] + wr)) // 2
        placed = [(b["left"], x0, (rh - hl) // 2, wl, hl)]
        y = (rh - hr) // 2
        for n, h in zip(b["rights"], hs):
            placed.append((n, x0 + wl + b["gap"], y, wr, h)); y += h + b["vgap"]
        return rh, placed

    def measure(self, b, cw):
        if b["k"] == "bar":
            return (self.BAR_H2 if "\n" in b["text"] else self.BAR_H1), self.GAP_BAR
        if b["k"] == "pic":
            return int(cw * b.get("frac", 1.0) / self._asp(b["name"])), self.GAP_FIG
        if b["k"] == "grid":
            return self._grid_layout(b, cw)[0], self.GAP_FIG
        if b["k"] == "split":
            return self._split_layout(b, cw)[0], self.GAP_FIG
        if b["k"] == "head":
            return self.txt_h([b["text"]], cw, self.HEAD_SZ, bold=True), 50000
        sa = int(b.get("space_after", Pt(0)).pt * 12700) if "space_after" in b else 0
        return self.txt_h(b["paras"], cw, b.get("size", self.BODY),
                          b.get("bullets", False), sa, b.get("bold", False)), 0

    def band_h(self, blocks, cw):
        hs, gs = zip(*[self.measure(b, cw) for b in blocks])
        return sum(hs) + sum(gs[:-1])

    # -------------------------------------------------------------- rendering
    def render(self, col, blocks, y0, y1, label=""):
        """Draw `blocks` into column `col` between y0 and y1, distributing the
        leftover height into the gaps that precede each section bar."""
        cw, cx = self.COL_WIDTH[col], self.COL_X[col]
        heights, gaps = zip(*[self.measure(b, cw) for b in blocks])
        starts = [i for i, b in enumerate(blocks)
                  if b["k"] == "bar" or (b["k"] == "head" and b.get("sec", True))]
        n_flex = max(0, len(starts) - 1)
        rigid = sum(heights) + sum(g for i, g in enumerate(gaps)
                                   if i + 1 not in starts and i < len(blocks) - 1)
        avail = y1 - y0
        slack = avail - rigid - n_flex * self.GAP_MIN
        flex = self.GAP_MIN + (max(0, slack) / n_flex if n_flex else 0)
        # A column with a single section has no bar gaps to absorb slack, so
        # spread it between the blocks instead of leaving a hole at the bottom.
        spread = 0
        if n_flex == 0 and slack > 0 and len(blocks) > 1:
            spread = slack / (len(blocks) - 1)
        note = ""
        if slack < 0:
            note = "   *** OVERFLOW — shrink a figure or move a block ***"
        elif avail and slack > 0.15 * avail:
            note = "   <- under-filled: enlarge the figures or add content"
        print(f"  [{label or col:12s}] content {rigid/CM:6.1f} cm | avail {avail/CM:6.1f} cm"
              f" | slack {slack/CM:6.1f} cm" + note)

        y = y0
        for i, b in enumerate(blocks):
            if b["k"] == "bar":
                sh = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(cx),
                                                 Emu(int(y)), Emu(cw), Emu(heights[i]))
                sh.fill.solid(); sh.fill.fore_color.rgb = self.C["accent"]
                sh.line.color.rgb = self.C["accent"]; sh.line.width = Pt(1)
                sh.shadow.inherit = False
                tf = sh.text_frame; tf.word_wrap = True
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for j, line in enumerate(b["text"].split("\n")):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.alignment = PP_ALIGN.CENTER
                    r = p.add_run(); r.text = line
                    self._fmt(r, self.BAR_SZ, True, self.C["white"])
            elif b["k"] == "pic":
                w = int(cw * b.get("frac", 1.0))
                self.slide.shapes.add_picture(f"{self.FIGS}/{b['name']}.png",
                                              Emu(cx + (cw - w) // 2), Emu(int(y)),
                                              Emu(w), Emu(heights[i]))
            elif b["k"] in ("grid", "split"):
                fn = self._grid_layout if b["k"] == "grid" else self._split_layout
                for item in fn(b, cw)[1]:
                    name, dx, dy, pw, ph = item[:5]
                    if len(item) > 5:
                        fx, fw, fh, flab = item[5:]
                        box = self.slide.shapes.add_shape(
                            MSO_SHAPE.ROUNDED_RECTANGLE, Emu(cx + fx), Emu(int(y)),
                            Emu(fw), Emu(fh))
                        box.adjustments[0] = 0.02
                        box.fill.solid(); box.fill.fore_color.rgb = self.C["frame_fill"]
                        box.line.color.rgb = self.C["frame_line"]
                        box.line.width = Pt(1.1); box.shadow.inherit = False
                        self.textbox(cx + fx + int(0.35 * CM), int(y) + int(0.10 * CM),
                                     fw - int(0.7 * CM), [flab], size=self.BODY,
                                     bold=True, color=self.C["accent"],
                                     align=PP_ALIGN.LEFT)
                    self.slide.shapes.add_picture(f"{self.FIGS}/{name}.png",
                                                  Emu(cx + dx), Emu(int(y) + dy),
                                                  Emu(pw), Emu(ph))
            elif b["k"] == "head":
                self.textbox(cx, int(y), cw, [b["text"]], size=self.HEAD_SZ,
                             bold=True, align=PP_ALIGN.LEFT)
            else:
                kw = {k: v for k, v in b.items() if k not in ("k", "paras")}
                self.textbox(cx, int(y), cw, b["paras"], **kw)
            y += heights[i]
            if i < len(blocks) - 1:
                y += flex if (i + 1 in starts) else gaps[i] + spread
        return y

    # ----------------------------------------------------------------- header
    def header(self, title, authors, affils, logo=None, qr=None, title_pt=None):
        s = self.W / (57.4 * CM)
        hdr = self.slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Emu(int(0.42 * CM)), Emu(int(0.32 * CM)),
            Emu(self.W - int(0.84 * CM)), Emu(self.HEADER_H - int(0.64 * CM)))
        hdr.adjustments[0] = 0.052
        hdr.fill.solid(); hdr.fill.fore_color.rgb = self.C["accent"]
        hdr.line.color.rgb = self.C["edge"]; hdr.line.width = Pt(0.75)
        hdr.shadow.inherit = False

        d = int(6.4 * CM * s)
        if logo:
            circ = self.slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(int(1.0 * CM)),
                                               Emu(int(0.95 * CM)), Emu(d), Emu(d))
            circ.fill.solid(); circ.fill.fore_color.rgb = self.C["white"]
            circ.line.fill.background(); circ.shadow.inherit = False
            p = int(d * 0.10)
            self.slide.shapes.add_picture(logo, Emu(int(1.0 * CM) + p),
                                          Emu(int(0.95 * CM) + p),
                                          Emu(d - 2 * p), Emu(d - 2 * p))
        if qr:
            self.slide.shapes.add_picture(qr, Emu(self.W - int(1.0 * CM) - d),
                                          Emu(int(0.95 * CM)), Emu(d), Emu(d))

        tpt = title_pt or Pt(round(34 * s, 1))
        self.textbox(0, int(1.15 * CM * s), self.W, title, size=tpt, bold=True,
                     color=self.C["white"], align=PP_ALIGN.CENTER)
        y = int(1.15 * CM * s) + self.txt_h(title, self.W, tpt, bold=True) + int(0.25 * CM)
        self.textbox(0, y, self.W, [authors], size=Pt(round(19 * s, 1)), bold=True,
                     color=self.C["white"], align=PP_ALIGN.CENTER)
        y += int(1.2 * CM * s)
        self.textbox(0, y, self.W, affils, size=Pt(round(14 * s, 1)), bold=True,
                     color=self.C["white"], align=PP_ALIGN.CENTER)

    # ---------------------------------------------------------------- compose
    def compose(self, bands):
        """bands: list of ("F", blocks) and at most one ("MID", left, right).
        Full-width bands take their natural height; MID absorbs the remainder."""
        top = self.HEADER_H + int(0.9 * CM)
        bot = self.H - self.BOT_PAD
        fixed = sum(self.band_h(b[1], self.FULL_W) for b in bands if b[0] == "F")
        mid = [i for i, b in enumerate(bands) if b[0] == "MID"]
        gaps = self.BAND_GAP * (len(bands) - 1)
        mid_h = (bot - top) - fixed - gaps
        if mid and mid_h < 0:
            print(f"  !! middle region would be {mid_h/CM:.1f} cm — trim the bands")

        y = top
        for i, band in enumerate(bands):
            if band[0] == "F":
                h = self.band_h(band[1], self.FULL_W)
                self.render("F", band[1], y, y + h, band[2] if len(band) > 2 else "")
                y += h
            else:
                self.render("L", band[1], y, y + mid_h, "mid-L")
                self.render("R", band[2], y, y + mid_h, "mid-R")
                y += mid_h
            y += self.BAND_GAP
        return self

    def save(self, path):
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        self.prs.save(path)
        print("saved", path)
        return path
